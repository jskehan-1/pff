"""
Loads PFF's preseason season-long fantasy projections (the raw stat lines,
not just a final point total) from the "pff" sheet of fantasyFB{YY}proj.xlsx
in the parent NFL folder, and converts them into a DK-scored per-game
baseline using our own dk_scoring.py formula.

This is the fallback used for anyone with zero real current-season (or
prior-season) game history in weekly_stats_*.csv -- mainly rookies, and
anyone who missed all of last season. It replaces DK's own AvgPointsPerGame
as the primary fallback since it's grounded in PFF's actual projected raw
stat line rather than a single opaque number, and DK's average is only
used if a player can't be found in this file at all.

The workbook lives outside the pff/ project folder (it's a general fantasy
football file, shared with the Flounder keeper league use case) -- this
looks for it by default, but the path can be overridden.
"""
import glob
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from dk_scoring import score_offense, score_dst  # noqa: E402

try:
    import openpyxl
except ImportError:
    openpyxl = None

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_SEARCH_DIR = PROJECT_ROOT.parent  # the NFL folder, one level up from pff/


def find_default_workbook(search_dir=None):
    """Picks the newest fantasyFB*proj.xlsx in the NFL folder, if any."""
    search_dir = Path(search_dir) if search_dir else DEFAULT_SEARCH_DIR
    candidates = sorted(
        glob.glob(str(search_dir / "fantasyFB*proj*.xlsx")),
        key=lambda p: Path(p).stat().st_mtime,
        reverse=True,
    )
    return Path(candidates[0]) if candidates else None


def _num(row, idx, key, default=0):
    v = row[idx[key]] if key in idx else None
    return v if isinstance(v, (int, float)) else default


def load_preseason_projections(workbook_path=None, sheet_name="pff"):
    """
    Returns dict: normalized key -> per-game DK-scored projection (float).
    Offense keyed by normalize_name(playerName) (import-compatible with
    build_projections.py's own normalize_name -- duplicated here minimally
    to avoid a circular import). DST keyed by "dst:" + canonical team code
    matching build_projections.py's dst_key() convention.
    """
    if openpyxl is None:
        print("[warn] openpyxl not installed -- can't read preseason projections "
              "xlsx. Install with: pip install openpyxl")
        return {}

    path = Path(workbook_path) if workbook_path else find_default_workbook()
    if not path or not path.exists():
        print(f"[warn] no preseason projections workbook found "
              f"(looked for fantasyFB*proj*.xlsx in {DEFAULT_SEARCH_DIR}) -- "
              f"rookies/no-history players will fall back to DK's own average.")
        return {}

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        print(f"[warn] {path.name} has no '{sheet_name}' sheet -- sheets found: "
              f"{wb.sheetnames}")
        return {}
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    idx = {h: i for i, h in enumerate(header) if h}

    # Local, minimal copies of the normalization helpers build_projections.py
    # uses, to avoid a circular import between the two modules.
    import re

    def normalize_name(name):
        name = str(name or "").strip().lower()
        name = re.sub(r"[.'\"]", "", name)
        name = re.sub(r"\b(jr|sr|ii|iii|iv)\b\.?", "", name)
        name = re.sub(r"\s+", " ", name).strip()
        return name

    TEAM_ALIASES = {
        "ari": "ARI", "arz": "ARI", "cardinals": "ARI",
        "atl": "ATL", "falcons": "ATL",
        "bal": "BAL", "blt": "BAL", "ravens": "BAL",
        "buf": "BUF", "bills": "BUF",
        "car": "CAR", "panthers": "CAR",
        "chi": "CHI", "bears": "CHI",
        "cin": "CIN", "bengals": "CIN",
        "cle": "CLE", "clv": "CLE", "browns": "CLE",
        "dal": "DAL", "cowboys": "DAL",
        "den": "DEN", "broncos": "DEN",
        "det": "DET", "lions": "DET",
        "gb": "GB", "gnb": "GB", "packers": "GB",
        "hou": "HOU", "hst": "HOU", "texans": "HOU",
        "ind": "IND", "colts": "IND",
        "jax": "JAX", "jac": "JAX", "jaguars": "JAX",
        "kc": "KC", "kan": "KC", "chiefs": "KC",
        "lv": "LV", "lvr": "LV", "oak": "LV", "raiders": "LV",
        "lac": "LAC", "sd": "LAC", "chargers": "LAC",
        "lar": "LAR", "la": "LAR", "stl": "LAR", "rams": "LAR",
        "mia": "MIA", "dolphins": "MIA",
        "min": "MIN", "vikings": "MIN",
        "ne": "NE", "nwe": "NE", "patriots": "NE",
        "no": "NO", "nor": "NO", "saints": "NO",
        "nyg": "NYG", "giants": "NYG",
        "nyj": "NYJ", "jets": "NYJ",
        "phi": "PHI", "eagles": "PHI",
        "pit": "PIT", "steelers": "PIT",
        "sea": "SEA", "seahawks": "SEA",
        "sf": "SF", "sfo": "SF", "49ers": "SF", "niners": "SF",
        "tb": "TB", "tam": "TB", "buccaneers": "TB", "bucs": "TB",
        "ten": "TEN", "titans": "TEN",
        "was": "WAS", "wsh": "WAS", "commanders": "WAS",
    }

    def dst_key(team_or_name):
        key = (team_or_name or "").strip().lower()
        return "dst:" + TEAM_ALIASES.get(key, key.upper())

    result = {}
    for row in rows[1:]:
        if not row or not row[idx.get("playerName", 0)]:
            continue
        position = str(_num(row, idx, "position", "") or row[idx["position"]] or "").lower()
        games = _num(row, idx, "games", 0) or 17

        if position == "dst":
            pts = score_dst(
                sacks=_num(row, idx, "dstSacks"), ints=_num(row, idx, "dstInt"),
                fumble_rec=_num(row, idx, "dstFumblesRecovered"),
                return_td=_num(row, idx, "dstTd"),
            )
            # Points-allowed tiers are expected GAME COUNTS per tier (confirmed:
            # they sum to `games`), so their DK bonus contributes as a season
            # total, not a single per-game value -- add it in before dividing.
            pa_bonus = (
                10 * _num(row, idx, "dstPts0") + 7 * _num(row, idx, "dstPts16") +
                4 * _num(row, idx, "dstPts713") + 1 * _num(row, idx, "dstPts1420") +
                0 * _num(row, idx, "dstPts2127") + -1 * _num(row, idx, "dstPts2834") +
                -4 * _num(row, idx, "dstPts35plus")
            )
            season_pts = pts + pa_bonus
            team_name = str(row[idx["playerName"]]).replace(" DST", "").strip()
            key = dst_key(team_name)
        else:
            season_pts = score_offense(
                pass_yds=_num(row, idx, "passYds"), pass_td=_num(row, idx, "passTd"),
                ints=_num(row, idx, "passInt"),
                rush_yds=_num(row, idx, "rushYds"), rush_td=_num(row, idx, "rushTd"),
                rec=_num(row, idx, "recvReceptions"), rec_yds=_num(row, idx, "recvYds"),
                rec_td=_num(row, idx, "recvTd"),
                fumbles_lost=_num(row, idx, "fumblesLost"), two_pt=_num(row, idx, "twoPt"),
                return_td=_num(row, idx, "returnTd"),
            )
            key = normalize_name(row[idx["playerName"]])

        result[key] = round(season_pts / games, 2) if games else 0.0

    return result


if __name__ == "__main__":
    proj = load_preseason_projections()
    print(f"Loaded {len(proj)} preseason per-game projections.")
    for name in ["josh allen", "christian mccaffrey", "dst:BUF"]:
        print(f"  {name}: {proj.get(name, 'NOT FOUND')}")
